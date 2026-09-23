"""Read partner ZIPs in memory. Never extract paths or execute spreadsheet content."""
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean
from openpyxl import load_workbook
from engine import detect_outliers, num

MONTHS = ['янв','фев','мар','апр','май','июн','июл','авг','сен','окт','ноя','дек']


def month_key(value):
    text = str(value).lower()
    year = re.search(r'20\d{2}', text)
    for i, prefix in enumerate(MONTHS):
        if year and text.startswith(prefix):
            return f'{year.group()}-{i+1:02d}'
    return None


def import_zip(payload, filename=''):
    items = {}
    sources = []
    transactions = defaultdict(list)
    types = Counter()
    seasonality = [1.0]*12
    with zipfile.ZipFile(io.BytesIO(payload), metadata_encoding='cp866') as archive:
        infos = [i for i in archive.infolist() if i.filename.lower().endswith('.xlsx')]
        if not infos or len(infos)>30 or sum(i.file_size for i in infos)>150_000_000:
            raise ValueError('ZIP должен содержать до 30 Excel-файлов, суммарно не более 150 МБ')
        supplier = 'Systeme Electric' if any('system' in i.filename.lower() for i in infos) else 'IEK'
        def item(code, name=None):
            code = str(code).strip()
            if code not in items:
                items[code] = dict(id=supplier+':'+code, supplier=supplier, code=code,
                    sku=code, name=name or code, unit='шт', category=None, sales={},
                    stock=None, stock_kind='missing', stock_date=None, stocks={}, inbound=[],
                    pack=0, minimum=0, anomalies=[], warnings=[], stockout_days={})
            if name:
                items[code]['name'] = str(name).strip()
            return items[code]
        # Monthly history first; shipment snapshots later override with actual free stock.
        def priority(info):
            low=info.filename.lower()
            return 3 if ('пути' in low or 'путь' in low) else 1
        for info in sorted(infos, key=priority):
            low = info.filename.lower()
            wb = load_workbook(io.BytesIO(archive.read(info)), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            source = dict(file=info.filename, sheet=ws.title, rows=ws.max_row, columns=ws.max_column, supplier=supplier)
            sources.append(source)
            rows = ws.iter_rows(values_only=True)
            if 'динамика' in low:
                for row in rows:
                    if not row[3] or row[3]=='Код':
                        continue
                    try:
                        dt = datetime.strptime(str(row[0]), '%d.%m.%Y %H:%M:%S')
                    except ValueError:
                        continue
                    doc = str(row[2] or '')
                    types[doc.split(' ')[0]] += 1
                    # Partner export: ordinary invoices are positive; negative corrections
                    # are already netted in monthly totals and are not demand spikes.
                    if doc.startswith('Расходная накладная') and num(row[7])>0:
                        transactions[str(row[3]).strip()].append(dict(date=dt.date().isoformat(),
                            document=str(row[1]),quantity=num(row[7]),warehouse=str(row[6])))
                source['kind']='transactions'
            elif 'ежемесячные продажи' in low:
                header=next(rows)
                code_ix=1
                columns={i:month_key(v) for i,v in enumerate(header) if month_key(v)}
                for row in rows:
                    if row[code_ix] is None:
                        continue
                    it=item(row[code_ix],row[0])
                    it['sales']={key:num(row[i]) for i,key in columns.items()}
                    if supplier=='Systeme Electric':
                        it['sku']=str(row[2] or it['code'])
                        # Monthly file's pack is diagnostic only; dedicated MOQ is authoritative.
                        it['monthly_pack']=num(row[3])
                source['kind']='sales'
            elif 'остатки' in low:
                header=next(rows)
                columns={i:month_key(v) for i,v in enumerate(header) if month_key(v)}
                for row in rows:
                    if row[2] is None:
                        continue
                    it=item(row[2],row[1] if supplier=='Systeme Electric' else row[0])
                    it['unit']=str(row[3] if supplier=='Systeme Electric' else row[1])
                    it['stocks']={key:(num(row[i]) if row[i] is not None else None) for i,key in columns.items()}
                    latest=max(columns.values())
                    it['stock']=it['stocks'].get(latest)
                    it['stock_date']=latest+'-01'
                    it['stock_kind']='monthly' if it['stock'] is not None else 'missing'
                source['kind']='stock'
            elif 'moq' in low:
                next(rows)
                for row in rows:
                    code=row[2] if supplier=='Systeme Electric' else row[1]
                    if not code:
                        continue
                    it=item(code,row[1] if supplier=='Systeme Electric' else row[3])
                    it['sku']=str(row[3] if supplier=='Systeme Electric' else row[2])
                    if supplier=='Systeme Electric':
                        it['pack']=num(row[4])
                    else:
                        # Minimum shipment is not necessarily order multiple.
                        it['minimum']=num(row[4])
                        it['pack']=1
                source['kind']='order_constraints'
            elif 'сезонность' in low:
                years=[]
                for row in rows:
                    if row[0] in (2024,2025):
                        values=[max(0,num(v)) for v in row[1:13]]
                        avg=mean(values)
                        if avg>0:
                            years.append([v/avg for v in values])
                if years:
                    seasonality=[mean(y[m] for y in years) for m in range(12)]
                source['kind']='seasonality'
            elif 'пути' in low or 'путь' in low:
                if supplier=='Systeme Electric':
                    next(rows); next(rows)
                    for row in rows:
                        if not row[2] or not row[3]:
                            continue
                        it=item(row[2],row[3]); it['sku']=str(row[1] or row[2])
                        it['category']=str(row[4]) if row[4] is not None else None
                        if row[51] is not None:
                            it.update(stock=num(row[51]),stock_kind='current',stock_date='2026-09-22')
                        if row[43] is not None:
                            it['growth_forecast']=num(row[43])
                        if num(row[54])>0:
                            it['inbound'].append(dict(eta='2026-09-24',quantity=num(row[54])))
                else:
                    header=next(rows)
                    dates={}
                    for i,v in enumerate(header[3:],3):
                        match=re.search(r'до (\d{2}\.\d{2}\.\d{4})',str(v))
                        if match:
                            dates[i]=datetime.strptime(match.group(1),'%d.%m.%Y').date().isoformat()
                    for row in rows:
                        if not row[0]:
                            continue
                        it=item(row[0],row[2]); it['sku']=str(row[1] or row[0])
                        it['unit_review']='БУХТАМИ' in str(row[2]).upper()
                        it['inbound']=[dict(eta=eta,quantity=num(row[i])) for i,eta in dates.items() if num(row[i])>0]
                source['kind']='inbound'
            wb.close()
        for code,it in items.items():
            it['seasonality']=seasonality
            it['anomalies']=detect_outliers(transactions.get(code,[]))
            if 'monthly_pack' in it and it['monthly_pack']!=it['pack'] and supplier=='Systeme Electric':
                it['warnings'].append('Кратность взята из MOQ; в месячной таблице другое значение')
            it['warnings'].append('В исходной динамике нет ID клиента; выбросы определены по накладным')
        return dict(items=list(items.values()),meta=dict(supplier=supplier,as_of='2026-09-23',
            source_date='2026-09-22',sources=sources,transactions=sum(map(len,transactions.values())),
            document_types=dict(types), notes=[
                'Пустые месячные продажи трактуются как нулевые продажи в форме 1С; пустые остатки остаются неизвестными.',
                'Начальные месячные остатки не подтверждают ежедневный stockout. Компенсация требует отдельного файла периодов.',
                'Коэффициенты сезонности поставщика получены по полным 2024 и 2025 годам; текущий неполный месяц не используется в обучении.',
                'Срок новой поставки и страховой запас — настраиваемые допущения, не условия из договоров.'
            ]))


def merge_datasets(datasets):
    items={i['id']:i for d in datasets for i in d['items']}
    counts={}
    for d in datasets:
        counts.update(d['meta'].get('transaction_counts',{}))
        if d['meta'].get('supplier'):
            counts[d['meta']['supplier']]=d['meta'].get('transactions',0)
    return dict(items=list(items.values()),meta=dict(as_of='2026-09-23',
        sources=[s for d in datasets for s in d['meta'].get('sources',[])],
        transaction_counts=counts,transactions=sum(counts.values()),
        notes=list(dict.fromkeys(n for d in datasets for n in d['meta'].get('notes',[])))))


def validate_normalized(data):
    if not isinstance(data,dict) or not isinstance(data.get('items'),list) or not 0<len(data['items'])<=10000:
        raise ValueError('JSON должен содержать items: массив из 1–10000 товаров')
    seen=set()
    for it in data['items']:
        for field in ('id','supplier','code','name','sales'):
            if field not in it:
                raise ValueError('Нет обязательного поля '+field)
        if it['id'] in seen:
            raise ValueError('Повтор идентификатора '+it['id'])
        seen.add(it['id'])
        if not isinstance(it['sales'],dict):
            raise ValueError('sales должен быть объектом месяц:количество')
        for key,value in it['sales'].items():
            datetime.strptime(key,'%Y-%m')
            if not isinstance(value,(int,float)) or num(value,-1)<0:
                raise ValueError('Продажи должны быть неотрицательными числами')
        if it.get('stock') is not None and (not isinstance(it['stock'],(int,float)) or num(it['stock'],-1)<0):
            raise ValueError('Остаток должен быть неотрицательным числом')
        it.setdefault('stock_kind','current' if it.get('stock') is not None else 'missing')
        it.setdefault('pack',1)
        for field in ('pack','minimum'):
            if field in it and (not isinstance(it[field],(int,float)) or num(it[field],-1)<0):
                raise ValueError(field+' должно быть неотрицательным числом')
        for key,value in it.get('stockout_days',{}).items():
            dt=datetime.strptime(key,'%Y-%m')
            import calendar
            if not isinstance(value,(int,float)) or not 0<=num(value,-1)<=calendar.monthrange(dt.year,dt.month)[1]:
                raise ValueError('Число дней stockout выходит за границы месяца')
        for p in it.get('inbound',[]):
            datetime.strptime(p['eta'],'%Y-%m-%d')
            if num(p.get('quantity'),-1)<0: raise ValueError('Некорректное количество в пути')
        if 'transactions' in it:
            it['anomalies']=detect_outliers(it.pop('transactions'))
    data.setdefault('meta',dict(notes=['Импортирован нормализованный JSON'],sources=[]))
    return data


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('archives',nargs='+')
    parser.add_argument('--output',default='data/dataset.json')
    args=parser.parse_args()
    ds=merge_datasets([import_zip(Path(p).read_bytes(),Path(p).name) for p in args.archives])
    out=Path(args.output); out.parent.mkdir(exist_ok=True,parents=True)
    out.write_text(json.dumps(ds,ensure_ascii=False),encoding='utf-8')
    print(json.dumps(dict(items=len(ds['items']),transactions=ds['meta']['transactions'],sources=len(ds['meta']['sources'])),ensure_ascii=False))
